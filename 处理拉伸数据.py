from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


STRAIN_UNITS = {"%", "mm/mm"}
STRESS_UNITS = {"MPa", "kN/mm\u00b2", "kN/mm2"}
DEFAULT_INPUT_DIR = Path(r"C:\Users\Administrator\Desktop\项目\实验数据\拉伸数据")
DEFAULT_OUTPUT_DIR = Path(r"C:\Users\Administrator\Desktop\项目\实验数据\拉伸曲线输出")
SLIP_THRESHOLD_PERCENT = 0.05
STRAIN_DECREASE_TOLERANCE_PERCENT = 0.05
TOE_TRIM_MIN_STRESS_MPA = 20.0
TOE_TRIM_RM_FRACTION = 0.02
TOE_TRIM_MAX_STRESS_MPA = 50.0
MIN_POINTS_AFTER_TRIM = 5
BRITTLE_AT_PERCENT = 1.0
BRITTLE_EXTENDED_AT_PERCENT = 1.5
BRITTLE_EXTENDED_MAX_STRAIN_PERCENT = 2.0
LOW_ELONGATION_AT_PERCENT = 3.0
YIELD_OFFSET_PERCENT = 0.2
ELASTIC_MIN_FIT_POINTS = 12
ELASTIC_MIN_STRAIN_SPAN_PERCENT = 0.015
ELASTIC_MIN_STRESS_SPAN_MPA = 25.0
SUSPECT_YIELD_MAX_RP_RM_RATIO = 0.60
SUSPECT_YIELD_MAX_DIRECT_FRACTURE_STRAIN_PERCENT = 5.0
SUSPECT_YIELD_MAX_POST_PEAK_STRAIN_PERCENT = 0.25
SUSPECT_YIELD_MAX_E_GPA = 80.0


def clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def to_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        if math.isnan(value) or math.isinf(value):
            return None
        return float(value)
    text = str(value).strip()
    if not text or text.startswith("#") or text in {"-----", "--", "-"}:
        return None
    text = text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def safe_rel(path: Path, root: Path) -> str:
    try:
        return str(path.relative_to(root))
    except ValueError:
        return str(path)


def safe_name(text: str) -> str:
    text = re.sub(r"[\\/:*?\"<>|]+", "_", text)
    text = re.sub(r"\s+", "_", text).strip("_")
    return text[:150] or "curve"


def wrap_legend_label(label: str, max_chars: int) -> list[str]:
    if max_chars <= 0 or len(label) <= max_chars:
        return [label]

    lines: list[str] = []
    remaining = label
    while len(remaining) > max_chars:
        cut = max_chars
        breakpoints = [
            remaining.rfind(" | ", 0, max_chars + 1),
            remaining.rfind("-", 0, max_chars + 1),
            remaining.rfind(" ", 0, max_chars + 1),
        ]
        usable = [idx for idx in breakpoints if idx >= max(1, max_chars // 3)]
        if usable:
            cut = max(usable)
            if remaining.startswith(" | ", cut):
                cut += 3
            else:
                cut += 1
        lines.append(remaining[:cut])
        remaining = remaining[cut:]
    if remaining:
        lines.append(remaining)
    return lines


def sample_status_key(sample_id: str) -> tuple[str, int]:
    text = sample_id.strip()
    match = re.match(r"^(.*?)[-\uff0d](\d+)$", text)
    if not match:
        return text, 0
    base = match.group(1)
    replicate = int(match.group(2))
    if re.fullmatch(r"C\d+", base, flags=re.IGNORECASE):
        return text, 0
    return base, replicate


def legend_rows_by_status(items: list[tuple[dict[str, Any], str, str]]) -> list[list[tuple[dict[str, Any], str, str]]]:
    groups: dict[str, list[tuple[int, tuple[dict[str, Any], str, str]]]] = {}
    order: dict[str, int] = {}
    for item in items:
        curve = item[0]
        base, replicate = sample_status_key(str(curve.get("sample_id") or curve.get("run_name") or ""))
        if base not in order:
            order[base] = len(order)
            groups[base] = []
        groups[base].append((replicate, item))

    paired_rows: list[list[tuple[dict[str, Any], str, str]]] = []
    single_items: list[tuple[dict[str, Any], str, str]] = []
    for base in sorted(order, key=order.get):
        group = [item for _, item in sorted(groups[base], key=lambda pair: pair[0])]
        while len(group) >= 2:
            paired_rows.append(group[:2])
            group = group[2:]
        single_items.extend(group)

    for idx in range(0, len(single_items), 2):
        paired_rows.append(single_items[idx : idx + 2])
    return paired_rows


def table_value(value: Any, digits: int) -> int | float | str:
    number = to_float(value)
    if number is None:
        return "-"
    rounded = round(number, digits)
    if digits == 0:
        return int(rounded)
    return rounded


def write_group_table(curves: list[dict[str, Any]], out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "性能"
    headers = ["样品", "Rp0.2 (MPa)", "Rm (MPa)", "均匀延伸率 (%)", "断裂延伸率 (%)", "E (GPa)"]
    ws.append(headers)
    for curve in curves:
        ws.append(
            [
                curve.get("sample_id") or curve.get("run_name") or "",
                table_value(curve.get("rp02_mpa"), 0),
                table_value(curve.get("rm_mpa"), 0),
                table_value(curve.get("uniform_elongation_percent"), 1),
                table_value(curve.get("at_percent"), 1),
                table_value(curve.get("e_gpa"), 0),
            ]
        )

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
    for col, width in enumerate([34, 14, 12, 16, 16, 10], start=1):
        ws.column_dimensions[chr(64 + col)].width = width
    ws.freeze_panes = "A2"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def file_plot_key_from_source(source_file: str, source_type: str) -> str:
    parts = [part for part in re.split(r"[\\/]+", source_file) if part]
    if source_type == "csv" and len(parts) >= 3 and parts[-2].endswith("_Exports"):
        return "\\".join(parts[:-2])
    return source_file


def parse_sample_id(sample_id: str) -> tuple[str, str, str]:
    sample_id = sample_id.strip()
    if not sample_id:
        return "\u672a\u586b\u5199", "", "\u672a\u586b\u5199"
    parts = [p for p in re.split(r"[-\uff0d]", sample_id) if p]
    first = parts[0] if parts else sample_id
    match = re.match(r"^(C\d+)", first, flags=re.IGNORECASE)
    material = match.group(1).upper() if match else first
    condition = "-".join(parts[1:-1]) if len(parts) >= 3 else ""
    # Group all samples with the same C-number into one figure.
    group = material
    return material, condition, group


def is_strain_unit(unit: str) -> bool:
    return unit in STRAIN_UNITS


def is_stress_unit(unit: str) -> bool:
    return unit in STRESS_UNITS or unit.startswith("kN/mm")


def header_key(text: str) -> str | None:
    packed = re.sub(r"\s+", "", text)
    if "\u540d\u79f0" in packed:
        return "name"
    if "\u8bd5\u6837\u6807\u8bc6" in packed:
        return "sample_id"
    if "Rp0.2" in packed or "\u89c4\u5b9a\u5851\u6027" in packed or "\u5c48\u670d\u5f3a\u5ea6" in packed:
        return "rp02_mpa"
    if "Rm" in packed or "\u6297\u62c9\u5f3a\u5ea6" in packed:
        return "rm_mpa"
    if "At" in packed or "\u65ad\u88c2\u603b\u4f38\u957f\u7387" in packed or "\u65ad\u88c2\u603b\u5ef6\u4f38\u7387" in packed or "\u65ad\u540e\u4f38\u957f\u7387" in packed:
        return "at_percent"
    if (
        packed in {"E", "E(GPa)", "E[GPa]"}
        or "\u5f39\u6027\u6a21\u91cf" in packed
        or "\u6768\u6c0f\u6a21\u91cf" in packed
        or "\u62c9\u4f38\u6a21\u91cf" in packed
    ):
        return "e_gpa"
    return None


def find_summary(ws: Any) -> tuple[int, dict[str, int]] | None:
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 80), max_col=min(ws.max_column or 1, 80), values_only=True),
        start=1,
    ):
        mapping: dict[str, int] = {}
        for col_idx, value in enumerate(row, start=1):
            key = header_key(clean(value))
            if key:
                mapping[key] = col_idx
        if {"name", "sample_id", "rp02_mpa", "rm_mpa", "at_percent"}.issubset(mapping):
            return row_idx, mapping
    return None


def load_run_info(ws: Any) -> dict[int, dict[str, Any]]:
    found = find_summary(ws)
    if not found:
        return {}
    header_row, mapping = found
    run_info: dict[int, dict[str, Any]] = {}
    for row_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row, values_only=True), start=header_row + 1):
        name = clean(row[mapping["name"] - 1] if mapping["name"] - 1 < len(row) else "")
        match = re.search(r"Test Run\s*(\d+)", name, flags=re.IGNORECASE)
        if not match:
            continue
        run_no = int(match.group(1))
        sample_id = clean(row[mapping["sample_id"] - 1] if mapping["sample_id"] - 1 < len(row) else "")
        material, condition, group = parse_sample_id(sample_id)
        run_info[run_no] = {
            "test_run": f"Test Run {run_no}",
            "sample_id": sample_id,
            "material": material,
            "condition": condition,
            "group": group,
            "rp02_mpa": to_float(row[mapping["rp02_mpa"] - 1] if mapping["rp02_mpa"] - 1 < len(row) else None),
            "rm_mpa": to_float(row[mapping["rm_mpa"] - 1] if mapping["rm_mpa"] - 1 < len(row) else None),
            "at_percent": to_float(row[mapping["at_percent"] - 1] if mapping["at_percent"] - 1 < len(row) else None),
            "e_gpa": to_float(row[mapping["e_gpa"] - 1] if "e_gpa" in mapping and mapping["e_gpa"] - 1 < len(row) else None),
        }
    return run_info


def find_curve_pairs(ws: Any) -> list[dict[str, Any]]:
    if (ws.max_row or 0) < 2:
        return []
    header = [clean(v).strip("[]") for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    units = [clean(v) for v in next(ws.iter_rows(min_row=2, max_row=2, values_only=True))]
    pairs: list[dict[str, Any]] = []
    for idx in range(len(units) - 1):
        left = units[idx]
        right = units[idx + 1]
        if is_strain_unit(left) and is_stress_unit(right):
            run_name = header[idx] or header[idx + 1] or f"Curve {len(pairs) + 1}"
            match = re.search(r"Test Run\s*(\d+)", run_name, flags=re.IGNORECASE)
            pairs.append(
                {
                    "start_idx": idx,
                    "run_name": run_name,
                    "run_no": int(match.group(1)) if match else None,
                    "strain_unit": left,
                    "stress_unit": right,
                }
            )
    return pairs


def downsample_xy(x: list[float], y: list[float], max_points: int = 2500) -> tuple[list[float], list[float]]:
    if len(x) <= max_points:
        return x, y
    step = max(1, len(x) // max_points)
    return x[::step], y[::step]


def linear_fit(points: list[tuple[int, float, float]]) -> dict[str, float] | None:
    if len(points) < 2:
        return None
    xs = [p[1] for p in points]
    ys = [p[2] for p in points]
    mean_x = sum(xs) / len(xs)
    mean_y = sum(ys) / len(ys)
    sxx = sum((x - mean_x) ** 2 for x in xs)
    if sxx <= 1e-12:
        return None
    sxy = sum((x - mean_x) * (y - mean_y) for x, y in zip(xs, ys))
    slope = sxy / sxx
    intercept = mean_y - slope * mean_x
    ss_tot = sum((y - mean_y) ** 2 for y in ys)
    ss_res = sum((y - (slope * x + intercept)) ** 2 for x, y in zip(xs, ys))
    r2 = 1.0 if ss_tot <= 1e-12 else max(0.0, min(1.0, 1.0 - ss_res / ss_tot))
    return {
        "slope_mpa_per_percent": slope,
        "intercept_mpa": intercept,
        "r2": r2,
        "point_count": float(len(points)),
        "strain_min_percent": min(xs),
        "strain_max_percent": max(xs),
        "stress_min_mpa": min(ys),
        "stress_max_mpa": max(ys),
        "first_index": float(min(p[0] for p in points)),
        "last_index": float(max(p[0] for p in points)),
    }


def fit_elastic_segment(strain_percent: list[float], stress_mpa: list[float], rm_mpa: float) -> dict[str, Any]:
    n = min(len(strain_percent), len(stress_mpa))
    if n < ELASTIC_MIN_FIT_POINTS:
        return {"fit": None, "flag": "NOT_ENOUGH_POINTS"}

    peak_idx = max(range(n), key=lambda idx: stress_mpa[idx])
    lower = max(10.0, rm_mpa * 0.02)
    upper_limit = min(max(150.0, rm_mpa * 0.45), rm_mpa * 0.75)
    base_points = [
        (idx, strain_percent[idx], stress_mpa[idx])
        for idx in range(0, peak_idx + 1)
        if strain_percent[idx] >= 0 and lower <= stress_mpa[idx] <= upper_limit
    ]
    if len(base_points) < ELASTIC_MIN_FIT_POINTS:
        base_points = [
            (idx, strain_percent[idx], stress_mpa[idx])
            for idx in range(0, peak_idx + 1)
            if strain_percent[idx] >= 0 and 10.0 <= stress_mpa[idx] <= min(rm_mpa * 0.8, 300.0)
        ]
    if len(base_points) < ELASTIC_MIN_FIT_POINTS:
        return {"fit": None, "flag": "NO_LINEAR_WINDOW"}

    candidates: list[dict[str, Any]] = []
    lower_fracs = [0.02, 0.03, 0.05, 0.08, 0.10, 0.12, 0.15]
    upper_fracs = [0.16, 0.20, 0.25, 0.30, 0.35, 0.40, 0.45]
    for lo_frac in lower_fracs:
        for hi_frac in upper_fracs:
            if hi_frac <= lo_frac:
                continue
            lo = max(10.0, rm_mpa * lo_frac)
            hi = min(upper_limit, max(lo + ELASTIC_MIN_STRESS_SPAN_MPA, rm_mpa * hi_frac))
            points = [p for p in base_points if lo <= p[2] <= hi]
            if len(points) < ELASTIC_MIN_FIT_POINTS:
                continue
            fit = linear_fit(points)
            if not fit:
                continue
            strain_span = fit["strain_max_percent"] - fit["strain_min_percent"]
            stress_span = fit["stress_max_mpa"] - fit["stress_min_mpa"]
            if strain_span < ELASTIC_MIN_STRAIN_SPAN_PERCENT or stress_span < ELASTIC_MIN_STRESS_SPAN_MPA:
                continue
            slope = fit["slope_mpa_per_percent"]
            if slope <= 0:
                continue
            e_gpa = slope / 10.0
            # A visually straight plastic-flow segment can have excellent R2 but an
            # impossible modulus. Do not use it as the elastic segment.
            if e_gpa < 60.0 or e_gpa > 350.0:
                continue
            modulus_penalty = 0.0
            if e_gpa < 50.0:
                modulus_penalty = min(0.08, (50.0 - e_gpa) / 500.0)
            elif e_gpa > 280.0:
                modulus_penalty = min(0.08, (e_gpa - 280.0) / 800.0)
            score = (
                fit["r2"]
                + min(0.01, math.log(len(points)) / 1000.0)
                + min(0.01, stress_span / max(rm_mpa, 1.0) / 20.0)
                - modulus_penalty
                - min(0.01, fit["strain_min_percent"] / 100.0)
            )
            candidates.append({"fit": fit, "score": score})

    if not candidates:
        points = base_points[: max(ELASTIC_MIN_FIT_POINTS, min(len(base_points), 80))]
        fit = linear_fit(points)
        if not fit:
            return {"fit": None, "flag": "FIT_FAILED"}
        flag = "FALLBACK_WINDOW"
    else:
        fit = max(candidates, key=lambda item: item["score"])["fit"]
        flag = "OK"

    e_gpa = fit["slope_mpa_per_percent"] / 10.0
    flags: list[str] = [flag]
    if int(fit["point_count"]) < ELASTIC_MIN_FIT_POINTS:
        flags.append("FEW_POINTS")
    if fit["r2"] < 0.98:
        flags.append("LOW_R2")
    if e_gpa < 50.0:
        flags.append("E_LOW")
    elif e_gpa > 280.0:
        flags.append("E_HIGH")
    return {"fit": fit, "flag": ";".join(f for f in flags if f and f != "OK") or "OK"}


def calculate_offset_yield(
    strain_percent: list[float],
    stress_mpa: list[float],
    fit: dict[str, float] | None,
) -> tuple[float | None, float | None, str]:
    if not fit:
        return None, None, "NO_ELASTIC_FIT"
    slope = fit["slope_mpa_per_percent"]
    intercept = fit["intercept_mpa"]
    if slope <= 0:
        return None, None, "BAD_ELASTIC_SLOPE"

    offset_values = [
        y - (slope * (x - YIELD_OFFSET_PERCENT) + intercept)
        for x, y in zip(strain_percent, stress_mpa)
    ]
    start_strain = max(YIELD_OFFSET_PERCENT, fit["strain_max_percent"])
    start_idx = 1
    for idx, x in enumerate(strain_percent):
        if x >= start_strain:
            start_idx = max(1, idx)
            break

    for idx in range(start_idx, len(offset_values)):
        prev_value = offset_values[idx - 1]
        curr_value = offset_values[idx]
        if prev_value >= 0 and curr_value <= 0:
            denom = prev_value - curr_value
            ratio = 0.0 if abs(denom) <= 1e-12 else prev_value / denom
            strain = strain_percent[idx - 1] + ratio * (strain_percent[idx] - strain_percent[idx - 1])
            stress = stress_mpa[idx - 1] + ratio * (stress_mpa[idx] - stress_mpa[idx - 1])
            return stress, strain, "OK"

    return None, None, "NO_OFFSET_INTERSECTION"


def should_suppress_offset_yield(
    *,
    rp02_mpa: float | None,
    rm_mpa: float | None,
    uniform_elongation_percent: float | None,
    at_percent: float | None,
    e_gpa: float | None,
) -> bool:
    if not all(value is not None for value in (rp02_mpa, rm_mpa, uniform_elongation_percent, at_percent, e_gpa)):
        return False
    if rm_mpa <= 0 or rp02_mpa <= 0:
        return False
    direct_fracture = abs(at_percent - uniform_elongation_percent) <= SUSPECT_YIELD_MAX_POST_PEAK_STRAIN_PERCENT
    low_direct_strain = uniform_elongation_percent <= SUSPECT_YIELD_MAX_DIRECT_FRACTURE_STRAIN_PERCENT
    low_yield_ratio = (rp02_mpa / rm_mpa) < SUSPECT_YIELD_MAX_RP_RM_RATIO
    low_modulus = e_gpa < SUSPECT_YIELD_MAX_E_GPA
    return direct_fracture and low_direct_strain and low_yield_ratio and low_modulus


def calculate_tensile_properties(strain_percent: list[float], stress_mpa: list[float]) -> dict[str, Any]:
    n = min(len(strain_percent), len(stress_mpa))
    if n < 3:
        return {
            "rp02_mpa": None,
            "yield_strain_percent": None,
            "rm_mpa": None,
            "uniform_elongation_percent": None,
            "at_percent": None,
            "e_gpa": None,
            "modulus_r2": None,
            "modulus_flag": "NOT_ENOUGH_POINTS",
            "yield_flag": "NOT_ENOUGH_POINTS",
        }
    peak_idx = max(range(n), key=lambda idx: stress_mpa[idx])
    rm_mpa = stress_mpa[peak_idx]
    uniform_elongation_percent = strain_percent[peak_idx]
    at_percent = strain_percent[n - 1]
    fit_result = fit_elastic_segment(strain_percent[:n], stress_mpa[:n], rm_mpa)
    fit = fit_result.get("fit")
    rp02_mpa, yield_strain_percent, yield_flag = calculate_offset_yield(strain_percent[:n], stress_mpa[:n], fit)
    e_gpa = fit["slope_mpa_per_percent"] / 10.0 if fit else None
    if should_suppress_offset_yield(
        rp02_mpa=rp02_mpa,
        rm_mpa=rm_mpa,
        uniform_elongation_percent=uniform_elongation_percent,
        at_percent=at_percent,
        e_gpa=e_gpa,
    ):
        rp02_mpa = None
        yield_strain_percent = None
        yield_flag = "SUSPECT_ELASTIC_WAVE"
    return {
        "rp02_mpa": rp02_mpa,
        "yield_strain_percent": yield_strain_percent,
        "rm_mpa": rm_mpa,
        "uniform_elongation_percent": uniform_elongation_percent,
        "at_percent": at_percent,
        "e_gpa": e_gpa,
        "modulus_r2": fit["r2"] if fit else None,
        "modulus_flag": fit_result.get("flag", ""),
        "yield_flag": yield_flag,
        "elastic_fit_start_percent": fit["strain_min_percent"] if fit else None,
        "elastic_fit_end_percent": fit["strain_max_percent"] if fit else None,
        "elastic_fit_points": int(fit["point_count"]) if fit else 0,
    }


def correct_extensometer_slip(strain_percent: list[float], threshold: float = SLIP_THRESHOLD_PERCENT) -> tuple[list[float], int, float]:
    if not strain_percent:
        return [], 0, 0.0

    corrected: list[float] = []
    offset = 0.0
    slip_count = 0
    slip_total = 0.0
    prev_raw = strain_percent[0]

    for idx, raw in enumerate(strain_percent):
        if idx > 0:
            jump = raw - prev_raw
            if abs(jump) > threshold:
                offset += jump
                slip_count += 1
                slip_total += abs(jump)
        corrected.append(raw - offset)
        prev_raw = raw

    return corrected, slip_count, slip_total


def trim_tail_fracture_jump(
    strain_percent: list[float],
    stress_mpa: list[float],
    threshold: float = SLIP_THRESHOLD_PERCENT,
) -> tuple[list[float], list[float], int]:
    """Drop the final one or two points when fracture causes an extensometer jump."""
    n = min(len(strain_percent), len(stress_mpa))
    if n < 3:
        return strain_percent[:n], stress_mpa[:n], 0

    start_idx = max(1, n - 2)
    for idx in range(start_idx, n):
        if abs(strain_percent[idx] - strain_percent[idx - 1]) > threshold:
            return strain_percent[:idx], stress_mpa[:idx], n - idx

    return strain_percent[:n], stress_mpa[:n], 0


def trim_after_strain_decrease(
    strain_percent: list[float],
    stress_mpa: list[float],
    tolerance: float = STRAIN_DECREASE_TOLERANCE_PERCENT,
) -> tuple[list[float], list[float], int]:
    """Keep the curve up to the largest strain before an unrecovered decrease."""
    n = min(len(strain_percent), len(stress_mpa))
    if n < 2:
        return strain_percent[:n], stress_mpa[:n], 0

    max_idx = 0
    max_strain = strain_percent[0]
    pending_keep: int | None = None
    pending_max = max_strain
    for idx in range(1, n):
        current = strain_percent[idx]
        if pending_keep is not None and current > pending_max + tolerance:
            pending_keep = None
        if current > max_strain:
            max_strain = current
            max_idx = idx
        if pending_keep is None and current < strain_percent[idx - 1] - tolerance:
            pending_keep = max_idx + 1
            pending_max = max_strain

    if pending_keep is not None:
        return strain_percent[:pending_keep], stress_mpa[:pending_keep], n - pending_keep

    return strain_percent[:n], stress_mpa[:n], 0


def trim_fracture_stress_drop(
    strain_percent: list[float],
    stress_mpa: list[float],
) -> tuple[list[float], list[float], int]:
    """Drop post-fracture instrument noise when stress collapses suddenly after the peak."""
    n = min(len(strain_percent), len(stress_mpa))
    if n < 5:
        return strain_percent[:n], stress_mpa[:n], 0

    peak_stress = max(stress_mpa[:n])
    if peak_stress <= 0:
        return strain_percent[:n], stress_mpa[:n], 0
    peak_idx = stress_mpa.index(peak_stress)
    low_level = max(5.0, peak_stress * 0.15)
    high_level = max(20.0, peak_stress * 0.50)
    min_drop = peak_stress * 0.35

    for idx in range(max(1, peak_idx + 1), n):
        prev_stress = stress_mpa[idx - 1]
        curr_stress = stress_mpa[idx]
        sudden_drop = prev_stress >= high_level and curr_stress <= low_level and (prev_stress - curr_stress) >= min_drop
        negative_after_peak = prev_stress >= high_level and curr_stress < 0
        if sudden_drop or negative_after_peak:
            return strain_percent[:idx], stress_mpa[:idx], n - idx

    return strain_percent[:n], stress_mpa[:n], 0


def toe_trim_threshold(rm_mpa: float | None) -> float:
    if rm_mpa is None or rm_mpa <= 0:
        return TOE_TRIM_MIN_STRESS_MPA
    return min(max(TOE_TRIM_MIN_STRESS_MPA, rm_mpa * TOE_TRIM_RM_FRACTION), TOE_TRIM_MAX_STRESS_MPA)


def trim_toe_and_zero(
    strain_percent: list[float],
    stress_mpa: list[float],
    rm_mpa: float | None,
) -> tuple[list[float], list[float], int, float, float, float]:
    """Remove the low-load seating section and set only the kept first strain to 0."""
    n = min(len(strain_percent), len(stress_mpa))
    if n < MIN_POINTS_AFTER_TRIM:
        return strain_percent[:n], stress_mpa[:n], 0, 0.0, 0.0, toe_trim_threshold(rm_mpa)

    threshold = toe_trim_threshold(rm_mpa)
    max_start = max(0, n - MIN_POINTS_AFTER_TRIM)
    start_idx = 0
    for idx, stress in enumerate(stress_mpa[: max_start + 1]):
        if stress >= threshold:
            start_idx = idx
            break

    zero_strain = strain_percent[start_idx]
    zero_stress = stress_mpa[start_idx]
    x = [value - zero_strain for value in strain_percent[start_idx:n]]
    y = stress_mpa[start_idx:n]
    return x, y, start_idx, zero_strain, zero_stress, threshold


def estimate_initial_modulus(
    strain_percent: list[float],
    stress_mpa: list[float],
    rm_mpa: float | None,
) -> tuple[float | None, float | None, str]:
    if len(strain_percent) < 5 or len(stress_mpa) < 5:
        return None, None, "NOT_ENOUGH_POINTS"

    if rm_mpa is not None and rm_mpa > 0:
        lower = min(50.0, max(10.0, rm_mpa * 0.05))
        upper = min(250.0, max(lower + 30.0, rm_mpa * 0.35))
    else:
        lower = 50.0
        upper = 250.0

    points = [
        (x, y)
        for x, y in zip(strain_percent, stress_mpa)
        if x > 0 and lower <= y <= upper
    ]
    if len(points) < 8:
        points = [
            (x, y)
            for x, y in zip(strain_percent, stress_mpa)
            if x > 0 and 20.0 <= y <= 200.0
        ]
    if len(points) < 5:
        return None, None, "NO_LINEAR_WINDOW"

    xs, ys = zip(*points)
    mean_x = sum(xs) / len(xs)
    mean_y = sum(ys) / len(ys)
    sxx = sum((x - mean_x) ** 2 for x in xs)
    if sxx <= 1e-12:
        return None, None, "NO_STRAIN_RANGE"
    sxy = sum((x - mean_x) * (y - mean_y) for x, y in points)
    slope_mpa_per_percent = sxy / sxx
    intercept = mean_y - slope_mpa_per_percent * mean_x
    ss_tot = sum((y - mean_y) ** 2 for y in ys)
    ss_res = sum((y - (slope_mpa_per_percent * x + intercept)) ** 2 for x, y in points)
    r2 = 1.0 if ss_tot <= 1e-12 else max(0.0, min(1.0, 1.0 - ss_res / ss_tot))
    e_gpa = slope_mpa_per_percent / 10.0

    flags: list[str] = []
    if e_gpa < 50.0:
        flags.append("E_LOW")
    elif e_gpa > 260.0:
        flags.append("E_HIGH")
    if r2 < 0.98:
        flags.append("LOW_R2")
    return e_gpa, r2, ";".join(flags)


def classify_fracture_mode(at_percent: float | None, max_strain_percent: float) -> str:
    reference_strain = at_percent if at_percent is not None and at_percent >= 0 else max_strain_percent
    if reference_strain <= BRITTLE_AT_PERCENT:
        return "脆断"
    if reference_strain <= BRITTLE_EXTENDED_AT_PERCENT and max_strain_percent <= BRITTLE_EXTENDED_MAX_STRAIN_PERCENT:
        return "脆断"
    if reference_strain <= LOW_ELONGATION_AT_PERCENT:
        return "低延伸率"
    return "塑性断裂"


def build_curve_record(
    *,
    source_file: str,
    source_type: str,
    sheet: str,
    run_name: str,
    run_no: int | None,
    sample_id: str,
    material: str,
    condition: str,
    group: str,
    strain_percent: list[float],
    stress_mpa: list[float],
    instrument: dict[str, Any] | None = None,
) -> dict[str, Any] | None:
    if len(strain_percent) < 5 or len(stress_mpa) < 5:
        return None
    instrument = instrument or {}
    raw_point_count = min(len(strain_percent), len(stress_mpa))
    rm_guess = max(stress_mpa[:raw_point_count]) if raw_point_count else None

    x, y, toe_trimmed_points, zero_strain_percent, zero_stress_mpa, toe_stress_threshold_mpa = trim_toe_and_zero(
        strain_percent,
        stress_mpa,
        rm_guess,
    )
    if len(x) < 5:
        return None
    x, y, fracture_drop_trimmed_points = trim_fracture_stress_drop(x, y)
    if len(x) < 5:
        return None
    x, y, tail_trimmed_points = trim_tail_fracture_jump(x, y)
    if len(x) < 5:
        return None
    x, y, strain_decrease_trimmed_points = trim_after_strain_decrease(x, y)
    if len(x) < 5:
        return None
    x, slip_count, slip_total_percent = correct_extensometer_slip(x)

    props = calculate_tensile_properties(x, y)
    max_strain_percent = max(x)
    fracture_mode = classify_fracture_mode(props.get("at_percent"), max_strain_percent)
    curve_note = ""
    if fracture_mode == "脆断":
        curve_note = "脆断样品，应变范围很小；Rp0.2 若无交点则为空。"
    elif props.get("yield_flag") != "OK":
        curve_note = "未找到稳定的0.2%偏移交点，屈服强度留空。"

    return {
        "source_file": source_file,
        "file_plot_key": file_plot_key_from_source(source_file, source_type),
        "source_type": source_type,
        "sheet": sheet,
        "run_name": run_name,
        "run_no": run_no,
        "sample_id": sample_id,
        "material": material,
        "condition": condition,
        "group": group,
        "rp02_mpa": props.get("rp02_mpa"),
        "yield_strain_percent": props.get("yield_strain_percent"),
        "rm_mpa": props.get("rm_mpa"),
        "uniform_elongation_percent": props.get("uniform_elongation_percent"),
        "at_percent": props.get("at_percent"),
        "e_gpa": props.get("e_gpa"),
        "instrument_rp02_mpa": instrument.get("rp02_mpa"),
        "instrument_rm_mpa": instrument.get("rm_mpa"),
        "instrument_at_percent": instrument.get("at_percent"),
        "instrument_e_gpa": instrument.get("e_gpa"),
        "raw_point_count": raw_point_count,
        "point_count": len(x),
        "max_stress_mpa": props.get("rm_mpa"),
        "max_strain_percent": max_strain_percent,
        "fracture_mode": fracture_mode,
        "curve_note": curve_note,
        "toe_trimmed_points": toe_trimmed_points,
        "zero_strain_percent": zero_strain_percent,
        "zero_stress_mpa": zero_stress_mpa,
        "toe_stress_threshold_mpa": toe_stress_threshold_mpa,
        "fracture_drop_trimmed_points": fracture_drop_trimmed_points,
        "slip_count": slip_count,
        "slip_total_percent": slip_total_percent,
        "tail_trimmed_points": tail_trimmed_points,
        "strain_decrease_trimmed_points": strain_decrease_trimmed_points,
        "estimated_e_gpa": props.get("e_gpa"),
        "modulus_r2": props.get("modulus_r2"),
        "modulus_flag": props.get("modulus_flag"),
        "yield_flag": props.get("yield_flag"),
        "elastic_fit_start_percent": props.get("elastic_fit_start_percent"),
        "elastic_fit_end_percent": props.get("elastic_fit_end_percent"),
        "elastic_fit_points": props.get("elastic_fit_points"),
        "x": x,
        "y": y,
    }


def load_curves(path: Path, root: Path, folder_run_info: dict[tuple[str, int], dict[str, Any]] | None = None) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    curves: list[dict[str, Any]] = []
    try:
        for ws in wb.worksheets:
            run_info = load_run_info(ws)
            pairs = find_curve_pairs(ws)
            if not pairs:
                continue
            xy = {i: ([], []) for i in range(len(pairs))}
            for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True):
                for pair_idx, pair in enumerate(pairs):
                    idx = pair["start_idx"]
                    strain = to_float(row[idx] if idx < len(row) else None)
                    stress = to_float(row[idx + 1] if idx + 1 < len(row) else None)
                    if strain is None or stress is None:
                        continue
                    if pair["strain_unit"] == "mm/mm":
                        strain *= 100.0
                    if pair["stress_unit"].startswith("kN/mm"):
                        stress *= 1000.0
                    xy[pair_idx][0].append(strain)
                    xy[pair_idx][1].append(stress)
            for pair_idx, pair in enumerate(pairs):
                x, y = xy[pair_idx]
                if len(x) < 5:
                    continue
                raw_point_count = len(x)
                parent_key = safe_rel(path.parent, root)
                run = run_info.get(pair["run_no"], {})
                if not run and pair["run_no"] is not None and folder_run_info is not None:
                    run = folder_run_info.get((parent_key, pair["run_no"]), {})
                sample_id = run.get("sample_id") or pair["run_name"]
                material, condition, group = parse_sample_id(sample_id)
                if run:
                    material = run.get("material", material)
                    condition = run.get("condition", condition)
                    group = run.get("group", group)
                record = build_curve_record(
                    source_file=safe_rel(path, root),
                    source_type="xlsx",
                    sheet=ws.title,
                    run_name=pair["run_name"],
                    run_no=pair["run_no"],
                    sample_id=sample_id,
                    material=material,
                    condition=condition,
                    group=group,
                    strain_percent=x,
                    stress_mpa=y,
                    instrument=run,
                )
                if record:
                    curves.append(record)
    finally:
        wb.close()
    return curves


def read_csv_rows(path: Path) -> list[list[str]]:
    for encoding in ("utf-8-sig", "gbk", "gb18030", "cp936"):
        try:
            with path.open("r", encoding=encoding, newline="") as f:
                return list(csv.reader(f))
        except UnicodeDecodeError:
            continue
    with path.open("r", encoding="utf-8", errors="ignore", newline="") as f:
        return list(csv.reader(f))


def load_csv_curve(path: Path, root: Path) -> list[dict[str, Any]]:
    rows = read_csv_rows(path)
    if not rows:
        return []

    sample_id = ""
    instrument: dict[str, Any] = {}
    run_no: int | None = None
    for row_idx, row in enumerate(rows):
        mapping: dict[str, int] = {}
        for col_idx, value in enumerate(row):
            key = header_key(clean(value))
            if key:
                mapping[key] = col_idx
        if "sample_id" not in mapping:
            continue
        for data_row in rows[row_idx + 1 :]:
            if mapping["sample_id"] >= len(data_row):
                continue
            candidate = clean(data_row[mapping["sample_id"]])
            if not candidate:
                continue
            sample_id = candidate
            run_no = int(to_float(data_row[0]) or 0) if data_row and to_float(data_row[0]) is not None else None
            instrument = {
                "sample_id": sample_id,
                "rp02_mpa": to_float(data_row[mapping["rp02_mpa"]]) if "rp02_mpa" in mapping and mapping["rp02_mpa"] < len(data_row) else None,
                "rm_mpa": to_float(data_row[mapping["rm_mpa"]]) if "rm_mpa" in mapping and mapping["rm_mpa"] < len(data_row) else None,
                "at_percent": to_float(data_row[mapping["at_percent"]]) if "at_percent" in mapping and mapping["at_percent"] < len(data_row) else None,
                "e_gpa": to_float(data_row[mapping["e_gpa"]]) if "e_gpa" in mapping and mapping["e_gpa"] < len(data_row) else None,
            }
            break
        if sample_id:
            break
    if not sample_id:
        sample_id = path.stem

    curve_header_idx: int | None = None
    strain_col: int | None = None
    stress_col: int | None = None
    for idx in range(1, len(rows)):
        units = [clean(v).strip("()[]") for v in rows[idx]]
        if "MPa" not in units or "%" not in units:
            continue
        headers = [clean(v) for v in rows[idx - 1]]
        if not any("\u65f6\u95f4" in header or "\u5e94\u53d8" in header or "\u5e94\u529b" in header for header in headers):
            continue
        stress_candidates = [
            col for col, unit in enumerate(units)
            if is_stress_unit(unit) or unit == "MPa"
        ]
        strain_candidates = [
            col for col, unit in enumerate(units)
            if is_strain_unit(unit)
        ]
        if not stress_candidates or not strain_candidates:
            continue
        stress_col = next((col for col in stress_candidates if col < len(headers) and "\u5e94\u529b" in headers[col]), stress_candidates[-1])
        strain_col = next(
            (
                col for col in strain_candidates
                if col < len(headers) and "\u62c9\u4f38\u5e94\u53d8" in headers[col]
            ),
            strain_candidates[0],
        )
        curve_header_idx = idx
        break

    if curve_header_idx is None or strain_col is None or stress_col is None:
        return []

    x: list[float] = []
    y: list[float] = []
    for row in rows[curve_header_idx + 1 :]:
        if max(strain_col, stress_col) >= len(row):
            continue
        strain = to_float(row[strain_col])
        stress = to_float(row[stress_col])
        if strain is None or stress is None:
            continue
        x.append(strain)
        y.append(stress)
    if len(x) < 5:
        return []

    material, condition, group = parse_sample_id(sample_id)
    record = build_curve_record(
        source_file=safe_rel(path, root),
        source_type="csv",
        sheet="",
        run_name=f"CSV {run_no or path.stem}",
        run_no=run_no,
        sample_id=sample_id,
        material=material,
        condition=condition,
        group=group,
        strain_percent=x,
        stress_mpa=y,
        instrument=instrument,
    )
    return [record] if record else []


def plot_curve_set(curves: list[dict[str, Any]], out_path: Path, title: str, max_legend: int = 60) -> None:
    from PIL import Image, ImageDraw, ImageFont

    plot_data: list[tuple[dict[str, Any], list[float], list[float]]] = []
    for curve in curves:
        valid = [
            (float(a), float(b))
            for a, b in zip(curve["x"], curve["y"])
            if isinstance(a, (int, float)) and isinstance(b, (int, float)) and a >= 0 and b >= 0
        ]
        if len(valid) < 2:
            continue
        x, y = zip(*valid)
        plot_data.append((curve, list(x), list(y)))

    x_values = [value for _, x, _ in plot_data for value in x]
    y_values = [value for _, _, y in plot_data for value in y]
    if not x_values or not y_values:
        return

    legend_count = min(len(plot_data), max_legend)
    legend_cols = 2 if legend_count > 1 else 1
    legend_rows = max(1, math.ceil(legend_count / legend_cols))
    width = 2400
    left, top, right, bottom = 140, 120, width - 90, 960
    legend_top = bottom + 95
    legend_row_h = 36
    height = legend_top + legend_rows * legend_row_h + 95
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)

    font_path = Path("C:/Windows/Fonts/msyh.ttc")
    if font_path.exists():
        title_font = ImageFont.truetype(str(font_path), 42)
        label_font = ImageFont.truetype(str(font_path), 28)
        tick_font = ImageFont.truetype(str(font_path), 22)
        legend_font = ImageFont.truetype(str(font_path), 21)
        legend_small_font = ImageFont.truetype(str(font_path), 18)
    else:
        title_font = label_font = tick_font = legend_font = legend_small_font = ImageFont.load_default()

    def fitted_legend_font(label: str, max_width: int):
        if not font_path.exists():
            return legend_small_font
        for size in range(18, 7, -1):
            font = ImageFont.truetype(str(font_path), size)
            bbox = draw.textbbox((0, 0), label, font=font)
            if bbox[2] - bbox[0] <= max_width:
                return font
        return ImageFont.truetype(str(font_path), 7)

    x_min = 0.0 if max(x_values) > 0 else min(x_values)
    x_max = max(x_values)
    y_min = 0.0 if max(y_values) > 0 else min(y_values)
    y_max = max(y_values)
    if x_max <= x_min:
        x_max = x_min + 1
    if y_max <= y_min:
        y_max = y_min + 1
    x_pad = (x_max - x_min) * 0.04
    y_pad = (y_max - y_min) * 0.08
    if x_min < 0:
        x_min -= x_pad
    x_max += x_pad
    if y_min < 0:
        y_min -= y_pad
    y_max += y_pad

    def sx(value: float) -> int:
        return int(left + (value - x_min) / (x_max - x_min) * (right - left))

    def sy(value: float) -> int:
        return int(bottom - (value - y_min) / (y_max - y_min) * (bottom - top))

    def x_tick_text(value: float) -> str:
        span = abs(x_max - x_min)
        if span < 0.1:
            return f"{value:.3f}"
        if span < 1:
            return f"{value:.2f}"
        return f"{value:.1f}"

    def fracture_summary_text() -> str:
        brittle_count = sum(1 for curve, _, _ in plot_data if curve.get("fracture_mode") == "脆断")
        if brittle_count <= 0:
            return ""
        return f"脆断：{brittle_count} 条"

    def fmt_perf(value: Any, digits: int = 0) -> str:
        number = to_float(value)
        if number is None:
            return "-"
        return f"{number:.{digits}f}"

    def legend_label(curve: dict[str, Any]) -> str:
        label = curve["sample_id"] or curve["run_name"]
        mode = curve.get("fracture_mode")
        if mode == "脆断":
            label = f"[{mode}] {label}"
        rp02 = fmt_perf(curve.get("rp02_mpa"), 0)
        rm = fmt_perf(curve.get("rm_mpa"), 0)
        e = fmt_perf(curve.get("e_gpa"), 0)
        au = fmt_perf(curve.get("uniform_elongation_percent"), 1)
        at = fmt_perf(curve.get("at_percent"), 1)
        return f"{label} | Rp0.2 {rp02} MPa | Rm {rm} MPa | 均匀 {au}% | 断裂 {at}% | E {e} GPa"

    palette = [
        "#1f77b4", "#ff7f0e", "#2ca02c", "#d62728", "#9467bd", "#8c564b",
        "#e377c2", "#7f7f7f", "#bcbd22", "#17becf", "#4e79a7", "#f28e2b",
        "#59a14f", "#e15759", "#76b7b2", "#edc949", "#af7aa1", "#ff9da7",
    ]

    draw.text((left, 35), title, fill="#111111", font=title_font)
    summary_text = fracture_summary_text()
    if summary_text:
        draw.text((left, 88), summary_text, fill="#555555", font=tick_font)

    for i in range(6):
        x = left + int((right - left) * i / 5)
        draw.line((x, top, x, bottom), fill="#E0E0E0", width=1)
        value = x_min + (x_max - x_min) * i / 5
        text = x_tick_text(value)
        draw.text((x - 28, bottom + 14), text, fill="#333333", font=tick_font)

    for i in range(6):
        y = bottom - int((bottom - top) * i / 5)
        draw.line((left, y, right, y), fill="#E0E0E0", width=1)
        value = y_min + (y_max - y_min) * i / 5
        text = f"{value:.0f}"
        draw.text((left - 92, y - 14), text, fill="#333333", font=tick_font)

    draw.rectangle((left, top, right, bottom), outline="#333333", width=2)
    draw.text(((left + right) // 2 - 58, bottom + 42), "\u5e94\u53d8 (%)", fill="#111111", font=label_font)
    draw.text((20, (top + bottom) // 2 - 20), "\u5e94\u529b (MPa)", fill="#111111", font=label_font)

    legend_items: list[tuple[dict[str, Any], str, str]] = []
    for idx, (curve, raw_x, raw_y) in enumerate(plot_data):
        x, y = downsample_xy(raw_x, raw_y, 2000)
        points = [(sx(float(a)), sy(float(b))) for a, b in zip(x, y)]
        if len(points) < 2:
            continue
        color = palette[idx % len(palette)]
        draw.line(points, fill=color, width=3, joint="curve")
        legend_items.append((curve, color, legend_label(curve)))

    draw.text((left, bottom + 76), "曲线标识与计算性能", fill="#111111", font=legend_font)
    col_w = (right - left) // legend_cols
    legend_rows_for_plot = legend_rows_by_status(legend_items[:max_legend])
    for row, row_items in enumerate(legend_rows_for_plot):
        for col, (_, color, label) in enumerate(row_items):
            legend_x = left + col * col_w
            legend_y = legend_top + row * legend_row_h
            draw.line((legend_x, legend_y + 13, legend_x + 42, legend_y + 13), fill=color, width=5)
            text_max_width = max(80, col_w - 58)
            item_font = fitted_legend_font(label, text_max_width)
            draw.text((legend_x + 52, legend_y), label, fill="#222222", font=item_font)
    if len(legend_items) > max_legend:
        draw.text((left, height - 58), f"另有 {len(legend_items) - max_legend} 条曲线未在图例中展开，详见 curve_index.csv", fill="#555555", font=legend_small_font)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    image.save(out_path)


def write_curve_index(path: Path, curves: list[dict[str, Any]], file_plot: str, group_plot: str) -> None:
    fields = [
        "source_file",
        "source_type",
        "sheet",
        "run_name",
        "sample_id",
        "group",
        "rp02_mpa",
        "yield_strain_percent",
        "rm_mpa",
        "uniform_elongation_percent",
        "at_percent",
        "e_gpa",
        "instrument_rp02_mpa",
        "instrument_rm_mpa",
        "instrument_at_percent",
        "instrument_e_gpa",
        "raw_point_count",
        "point_count",
        "max_stress_mpa",
        "max_strain_percent",
        "fracture_mode",
        "curve_note",
        "toe_trimmed_points",
        "zero_strain_percent",
        "zero_stress_mpa",
        "toe_stress_threshold_mpa",
        "fracture_drop_trimmed_points",
        "slip_count",
        "slip_total_percent",
        "tail_trimmed_points",
        "strain_decrease_trimmed_points",
        "estimated_e_gpa",
        "modulus_r2",
        "modulus_flag",
        "yield_flag",
        "elastic_fit_start_percent",
        "elastic_fit_end_percent",
        "elastic_fit_points",
        "file_plot",
        "group_plot",
    ]
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for curve in curves:
            row = {key: curve.get(key, "") for key in fields}
            row["file_plot"] = file_plot
            row["group_plot"] = group_plot
            writer.writerow(row)


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(description="Plot tensile stress-strain curves and calculate properties from raw Excel/CSV curves.")
    parser.add_argument("--input", default=str(DEFAULT_INPUT_DIR), help=f"Input folder. Default: {DEFAULT_INPUT_DIR}")
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT_DIR), help=f"Output folder. Default: {DEFAULT_OUTPUT_DIR}")
    args = parser.parse_args()

    input_root = Path(args.input)
    output_root = Path(args.output)
    if not input_root.exists():
        raise SystemExit(f"Input folder does not exist: {input_root}")
    by_file_dir = output_root / "by_file"
    by_group_dir = output_root / "by_group"
    output_root.mkdir(parents=True, exist_ok=True)
    by_file_dir.mkdir(parents=True, exist_ok=True)
    by_group_dir.mkdir(parents=True, exist_ok=True)
    for folder in (by_file_dir, by_group_dir):
        for old_plot in folder.glob("*.png"):
            old_plot.unlink()
    for old_table in by_group_dir.glob("*.xlsx"):
        old_table.unlink()

    files = sorted(p for p in input_root.rglob("*.xlsx") if not p.name.startswith("~$") and output_root not in p.parents)
    csv_files = sorted(
        p for p in input_root.rglob("*.csv")
        if output_root not in p.parents and p.parent.name.endswith("_Exports")
    )
    folder_run_info: dict[tuple[str, int], dict[str, Any]] = {}
    for file in files:
        try:
            wb = load_workbook(file, read_only=True, data_only=True)
            try:
                parent_key = safe_rel(file.parent, input_root)
                for ws in wb.worksheets:
                    for run_no, info in load_run_info(ws).items():
                        if info.get("sample_id"):
                            folder_run_info[(parent_key, run_no)] = info
            finally:
                wb.close()
        except Exception:
            continue

    all_curves: list[dict[str, Any]] = []
    file_plot_map: dict[str, str] = {}
    errors: list[dict[str, str]] = []

    for file in files:
        try:
            curves = load_curves(file, input_root, folder_run_info)
        except Exception as exc:
            errors.append({"source_file": safe_rel(file, input_root), "error": f"{type(exc).__name__}: {exc}"})
            continue
        if not curves:
            continue
        all_curves.extend(curves)

    for file in csv_files:
        try:
            curves = load_csv_curve(file, input_root)
        except Exception as exc:
            errors.append({"source_file": safe_rel(file, input_root), "error": f"{type(exc).__name__}: {exc}"})
            continue
        if not curves:
            continue
        all_curves.extend(curves)

    file_grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for curve in all_curves:
        file_grouped[curve.get("file_plot_key") or curve["source_file"]].append(curve)
    for file_key, curves in sorted(file_grouped.items()):
        out = by_file_dir / f"{safe_name(file_key)}_all_curves.png"
        plot_curve_set(curves, out, f"{file_key} - \u62c9\u4f38\u66f2\u7ebf")
        file_plot_map[file_key] = str(out)

    group_plot_map: dict[str, str] = {}
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for curve in all_curves:
        grouped[curve["group"]].append(curve)
    for group, curves in sorted(grouped.items()):
        out = by_group_dir / f"{safe_name(group)}_curves.png"
        table_out = out.with_suffix(".xlsx")
        plot_curve_set(curves, out, f"{group} - \u62c9\u4f38\u66f2\u7ebf")
        write_group_table(curves, table_out)
        group_plot_map[group] = str(out)

    index_path = output_root / "curve_index.csv"
    fields = [
        "source_file",
        "file_plot_key",
        "source_type",
        "sheet",
        "run_name",
        "sample_id",
        "group",
        "rp02_mpa",
        "yield_strain_percent",
        "rm_mpa",
        "uniform_elongation_percent",
        "at_percent",
        "e_gpa",
        "instrument_rp02_mpa",
        "instrument_rm_mpa",
        "instrument_at_percent",
        "instrument_e_gpa",
        "raw_point_count",
        "point_count",
        "max_stress_mpa",
        "max_strain_percent",
        "fracture_mode",
        "curve_note",
        "toe_trimmed_points",
        "zero_strain_percent",
        "zero_stress_mpa",
        "toe_stress_threshold_mpa",
        "fracture_drop_trimmed_points",
        "slip_count",
        "slip_total_percent",
        "tail_trimmed_points",
        "strain_decrease_trimmed_points",
        "estimated_e_gpa",
        "modulus_r2",
        "modulus_flag",
        "yield_flag",
        "elastic_fit_start_percent",
        "elastic_fit_end_percent",
        "elastic_fit_points",
        "file_plot",
        "group_plot",
    ]
    with index_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for curve in all_curves:
            row = {key: curve.get(key, "") for key in fields}
            row["file_plot"] = file_plot_map.get(curve.get("file_plot_key") or curve["source_file"], "")
            row["group_plot"] = group_plot_map.get(curve["group"], "")
            writer.writerow(row)
        for err in errors:
            writer.writerow({"source_file": err["source_file"], "sample_id": "ERROR", "group": err["error"]})

    print(
        {
            "xlsx_file_count": len(files),
            "csv_file_count": len(csv_files),
            "curve_count": len(all_curves),
            "file_plot_count": len(file_plot_map),
            "group_plot_count": len(group_plot_map),
            "error_count": len(errors),
            "output": str(output_root),
        }
    )


if __name__ == "__main__":
    main()
