"""Session-local, bounded upload adapter around the unchanged tensile CLI.

There is deliberately no Streamlit dependency or result cache in this module.
Each call owns one temporary directory and one child process. Only result bytes
escape that directory; callers decide how long to keep them in their session.
"""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path, PurePosixPath
import csv
import json
import math
import os
import re
import stat
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
import struct
import zlib

from openpyxl import Workbook, load_workbook


MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TOTAL_BYTES = 60 * 1024 * 1024
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_DATA_FILES = 30
_MAX_ARCHIVE_ENTRIES = 2000
_ZIP_IO_CHUNK = 64 * 1024
_LOCAL_FILE_HEADER = struct.Struct("<4s5H3L2H")
_CLI_SCRIPT = Path(__file__).with_name("处理拉伸数据.py")
_TEXT_FIELDS = {
    "source_file", "file_plot_key", "source_type", "sheet", "run_name",
    "sample_id", "group", "fracture_mode", "curve_note", "modulus_flag",
    "yield_flag", "file_plot", "group_plot",
}


@dataclass(frozen=True)
class Upload:
    name: str
    data: bytes


@dataclass
class ProcessingResult:
    rows: list[dict[str, str]]
    files: dict[str, bytes]
    warnings: list[str]
    errors: list[dict[str, str]]
    uploaded_count: int
    archive: bytes


class ProcessingError(Exception):
    """A Chinese, user-displayable validation or processing failure."""


@dataclass
class _Budget:
    expanded: int = 0

    def add(self, amount: int) -> None:
        self.expanded += amount
        if amount < 0 or self.expanded > MAX_EXPANDED_BYTES:
            raise ProcessingError("文件展开体积超过限制（含 XLSX 内部解压内容）。")


def _path_parts(name: str, *, flat: bool = False) -> tuple[str, ...]:
    """Validate before normalizing, including Windows aliases on Linux hosts."""
    name = unicodedata.normalize("NFC", name)
    portable = name.replace("\\", "/")
    parts = portable.split("/")
    if portable.endswith("/") and not flat:
        parts.pop()
    if (not parts or (flat and len(parts) != 1) or
            any(not part or part in {".", ".."} for part in parts)):
        raise ProcessingError(f"文件名或压缩包路径不安全：{name!r}")
    for part in parts:
        if (part.endswith((" ", ".")) or len(part) > 180 or
                any(ord(ch) < 32 or ch in ':*?"<>|' for ch in part) or
                re.fullmatch(r"(?:CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\..*)?", part, re.I)):
            raise ProcessingError(f"文件名或压缩包路径不安全：{name!r}")
    return tuple(parts)


def _archive_members(archive: zipfile.ZipFile) -> list[zipfile.ZipInfo]:
    members = archive.infolist()
    if len(members) > _MAX_ARCHIVE_ENTRIES:
        raise ProcessingError("压缩包条目数量过多。")
    seen: set[str] = set()
    for info in members:
        parts = _path_parts(info.orig_filename)
        key = "/".join(parts).casefold()
        if key in seen:
            raise ProcessingError(f"压缩包中存在同名冲突：{info.filename}")
        seen.add(key)
        if stat.S_ISLNK(info.external_attr >> 16):
            raise ProcessingError("压缩包不能包含符号链接。")
        if info.flag_bits & 1:
            raise ProcessingError("不支持加密压缩包，请先解密后上传。")
    return members


def _read_member_bounded(raw_archive: bytes, info: zipfile.ZipInfo,
                         budget: _Budget, *, collect: bool) -> bytes:
    """Decompress one member with a hard output cap independent of ZIP metadata."""
    offset = info.header_offset
    if offset < 0 or offset + _LOCAL_FILE_HEADER.size > len(raw_archive):
        raise ProcessingError("ZIP 压缩包结构损坏，请重新压缩后上传。")
    try:
        fields = _LOCAL_FILE_HEADER.unpack_from(raw_archive, offset)
    except struct.error as exc:
        raise ProcessingError("ZIP 压缩包结构损坏，请重新压缩后上传。") from exc
    signature, _, flags, method, _, _, _, _, _, name_length, extra_length = fields
    if (signature != b"PK\x03\x04" or flags & 1 or method != info.compress_type or
            method not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}):
        raise ProcessingError("ZIP 使用了不支持或损坏的压缩格式，请改用标准 ZIP（存储或 Deflate）。")
    data_start = offset + _LOCAL_FILE_HEADER.size + name_length + extra_length
    data_end = data_start + info.compress_size
    if info.compress_size < 0 or data_start < 0 or data_end > len(raw_archive):
        raise ProcessingError("ZIP 压缩包结构损坏，请重新压缩后上传。")

    chunks: list[bytes] = []
    actual_size = 0
    checksum = 0

    def emit(chunk: bytes) -> None:
        nonlocal actual_size, checksum
        if not chunk:
            return
        budget.add(len(chunk))
        actual_size += len(chunk)
        checksum = zlib.crc32(chunk, checksum)
        if collect:
            chunks.append(chunk)

    try:
        compressed = memoryview(raw_archive)[data_start:data_end]
        if method == zipfile.ZIP_STORED:
            for start in range(0, len(compressed), _ZIP_IO_CHUNK):
                emit(bytes(compressed[start:start + _ZIP_IO_CHUNK]))
        else:
            inflater = zlib.decompressobj(-15)
            for start in range(0, len(compressed), _ZIP_IO_CHUNK):
                pending = bytes(compressed[start:start + _ZIP_IO_CHUNK])
                while True:
                    remaining = MAX_EXPANDED_BYTES - budget.expanded
                    output_limit = min(_ZIP_IO_CHUNK, remaining + 1)
                    piece = inflater.decompress(pending, output_limit)
                    pending = inflater.unconsumed_tail
                    emit(piece)
                    if inflater.unused_data:
                        raise ProcessingError("ZIP 压缩包结构损坏，请重新压缩后上传。")
                    if pending:
                        continue
                    if inflater.eof or len(piece) < output_limit:
                        break
                    # zlib can consume all input while retaining output. An
                    # empty-input pass drains that output under the same cap.
            if not inflater.eof or inflater.unused_data or inflater.unconsumed_tail:
                raise ProcessingError("ZIP 压缩包结构损坏，请重新压缩后上传。")
    except zlib.error as exc:
        raise ProcessingError("ZIP 压缩数据损坏，请重新压缩后上传。") from exc

    if actual_size != info.file_size or checksum & 0xFFFFFFFF != info.CRC:
        raise ProcessingError("ZIP 压缩包大小或校验信息不一致，请重新压缩后上传。")
    return b"".join(chunks)


def _check_xlsx(data: bytes, budget: _Budget) -> None:
    # A malformed workbook is left for the original CLI so other valid curves
    # can still succeed. Valid ZIP containers must pass the expansion guard.
    try:
        archive = zipfile.ZipFile(BytesIO(data))
    except zipfile.BadZipFile:
        return
    with archive:
        for info in _archive_members(archive):
            if not info.is_dir():
                _read_member_bounded(data, info, budget, collect=False)


def _adapt_path(parts: tuple[str, ...]) -> PurePosixPath:
    path = PurePosixPath(*parts)
    path = path.with_suffix(path.suffix.lower())
    if path.suffix == ".csv" and not path.parent.name.endswith("_Exports"):
        path = path.parent / "_Exports" / path.name
    return path


def _stage_uploads(uploads: list[Upload], input_root: Path) -> tuple[list[str], list[str]]:
    if not uploads:
        raise ProcessingError("请先上传 XLSX、CSV 或 ZIP 文件。")
    budget = _Budget()
    total_bytes = 0
    top_names: set[str] = set()
    staged: dict[str, bytes] = {}
    warnings: list[str] = []

    def add_file(parts: tuple[str, ...], data: bytes) -> None:
        path = _adapt_path(parts)
        key = path.as_posix().casefold()
        for existing in staged:
            old = existing.casefold()
            if key == old or key.startswith(old + "/") or old.startswith(key + "/"):
                raise ProcessingError(f"文件路径存在同名冲突，不能覆盖：{path}")
        if len(staged) >= MAX_DATA_FILES:
            raise ProcessingError(f"数据文件数量超过限制（最多 {MAX_DATA_FILES} 个）。")
        if path.suffix == ".xlsx":
            _check_xlsx(data, budget)
        staged[path.as_posix()] = data

    for index, upload in enumerate(uploads, start=1):
        if not isinstance(upload.name, str) or not isinstance(upload.data, bytes):
            raise ProcessingError("上传内容无效，请重新选择文件。")
        parts = _path_parts(upload.name, flat=True)
        name = parts[0]
        if name.casefold() in top_names:
            raise ProcessingError(f"上传文件存在同名冲突：{name}")
        top_names.add(name.casefold())
        if len(upload.data) > MAX_FILE_BYTES:
            raise ProcessingError(f"文件大小超过限制：{name}（每个文件最多 {MAX_FILE_BYTES // 1024 // 1024} MB）。")
        total_bytes += len(upload.data)
        if total_bytes > MAX_TOTAL_BYTES:
            raise ProcessingError(f"上传总大小超过限制（最多 {MAX_TOTAL_BYTES // 1024 // 1024} MB）。")
        suffix = PurePosixPath(name).suffix.lower()
        if suffix not in {".xlsx", ".csv", ".zip"}:
            raise ProcessingError(f"不支持的文件格式：{name}；请使用 XLSX、CSV 或 ZIP。")
        if suffix != ".zip":
            budget.add(len(upload.data))
            add_file(parts, upload.data)
            continue
        try:
            with zipfile.ZipFile(BytesIO(upload.data)) as archive:
                members = _archive_members(archive)
                archive_root = f"zip_{index:02d}_{PurePosixPath(name).stem}"
                for info in members:
                    member_parts = _path_parts(info.orig_filename)
                    if info.is_dir():
                        continue
                    member_suffix = PurePosixPath(member_parts[-1]).suffix.lower()
                    if member_suffix == ".zip":
                        raise ProcessingError("ZIP 中不能包含嵌套 ZIP 压缩包。")
                    if member_suffix not in {".xlsx", ".csv"}:
                        warnings.append(f"已跳过非数据文件：{name}/{info.filename}")
                        continue
                    if info.file_size > MAX_EXPANDED_BYTES:
                        raise ProcessingError("压缩包展开体积超过限制。")
                    data = _read_member_bounded(upload.data, info, budget, collect=True)
                    add_file((archive_root, *member_parts), data)
        except ProcessingError:
            raise
        except (zipfile.BadZipFile, RuntimeError, OSError, NotImplementedError) as exc:
            raise ProcessingError(f"ZIP 压缩包无法读取：{name}；请检查文件是否完整。") from exc
    if not staged:
        raise ProcessingError("上传文件中没有可处理的 XLSX 或 CSV 数据文件。")
    for relative, data in staged.items():
        target = input_root.joinpath(*PurePosixPath(relative).parts)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(data)
    return list(staged), warnings


def _safe_text(value: str) -> str:
    # Quoting CSV alone does not stop spreadsheet formula interpretation.
    if value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
        return "'" + value
    return value


def _clean_workbook(path: Path) -> None:
    wb = load_workbook(path, data_only=False)
    try:
        for ws in wb.worksheets:
            for row in ws:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.value = _safe_text(cell.value)
        wb.save(path)
    finally:
        wb.close()


def _collect_results(output_root: Path, sources: list[str], warnings: list[str],
                     uploaded_count: int, temporary_root: Path) -> ProcessingResult:
    index_path = output_root / "curve_index.csv"
    if not index_path.is_file():
        raise ProcessingError("处理进程未生成结果索引，请检查上传文件格式。")
    with index_path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        fields = reader.fieldnames or []
        raw_rows = list(reader)
    errors: list[dict[str, str]] = []
    rows: list[dict[str, str]] = []
    produced: set[str] = set()
    failed: set[str] = set()
    plot_owners: dict[tuple[str, str], str] = {}
    for raw in raw_rows:
        row = {key: value or "" for key, value in raw.items() if key is not None}
        source = row.get("source_file", "").replace("\\", "/")
        if row.get("sample_id") == "ERROR" and not row.get("file_plot"):
            detail = row.get("group", "").replace(str(temporary_root), "[临时目录]")
            detail = detail.replace(temporary_root.as_posix(), "[临时目录]")
            errors.append({"source_file": _safe_text(source), "error": "读取失败：" + detail})
            failed.add(source)
            continue
        produced.add(source)
        row["source_file"] = source
        row["file_plot_key"] = row.get("file_plot_key", "").replace("\\", "/")
        for field in ("file_plot", "group_plot"):
            raw_path = row.get(field, "")
            try:
                relative = Path(raw_path).relative_to(output_root)
            except ValueError as exc:
                raise ProcessingError("结果图片路径无效，已停止打包。") from exc
            if not (output_root / relative).is_file():
                raise ProcessingError("结果图片缺失，已停止打包。")
            row[field] = relative.as_posix()
            owner = row.get("group", "") if field == "group_plot" else row.get("file_plot_key") or source
            plot_key = (field, row[field].casefold())
            if plot_key in plot_owners and plot_owners[plot_key] != owner:
                raise ProcessingError("结果文件名存在冲突，请调整上传文件名或样品标识后重试。")
            plot_owners[plot_key] = owner
        rows.append({key: _safe_text(value) if key in _TEXT_FIELDS else value for key, value in row.items()})
    for source in sources:
        if source not in produced and source not in failed:
            warnings.append(f"未产生曲线：{source}")
    if not rows:
        details = [e["source_file"] + "：" + e["error"] for e in errors] + warnings
        raise ProcessingError("没有识别到有效拉伸曲线。" + ("\n" + "\n".join(details[:8]) if details else ""))
    with index_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    for path in output_root.rglob("*.xlsx"):
        _clean_workbook(path)
    report = {
        "uploaded_count": uploaded_count, "data_file_count": len(sources),
        "curve_count": len(rows), "warnings": warnings, "errors": errors,
        "note": "仅包含处理结果；合成演示不代表真实实验精度。",
    }
    (output_root / "processing_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    files = {path.relative_to(output_root).as_posix(): path.read_bytes()
             for path in sorted(output_root.rglob("*")) if path.is_file()}
    archive_bytes = BytesIO()
    with zipfile.ZipFile(archive_bytes, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return ProcessingResult(rows, files, warnings, errors, uploaded_count, archive_bytes.getvalue())


def process_uploads(uploads: list[Upload], *, timeout_seconds: float = 120) -> ProcessingResult:
    """Validate uploads, run the CLI in isolation and return output-only bytes."""
    if not isinstance(timeout_seconds, (int, float)) or not math.isfinite(timeout_seconds) or timeout_seconds <= 0:
        raise ProcessingError("处理超时时间必须为大于零的有限秒数。")
    try:
        with tempfile.TemporaryDirectory(prefix="tensile_web_") as temporary:
            root = Path(temporary)
            input_root, output_root = root / "input", root / "output"
            input_root.mkdir()
            sources, warnings = _stage_uploads(uploads, input_root)
            command = [sys.executable, str(_CLI_SCRIPT), "--input", str(input_root), "--output", str(output_root)]
            environment = os.environ.copy()
            environment["PYTHONIOENCODING"] = "utf-8"
            try:
                completed = subprocess.run(command, cwd=root, env=environment, capture_output=True,
                                           timeout=timeout_seconds, check=False)
            except subprocess.TimeoutExpired as exc:
                raise ProcessingError(f"处理超时（{timeout_seconds:g} 秒），请减少文件数量后重试。") from exc
            if completed.returncode:
                # Keep tracebacks/server paths out of the user-facing response.
                raise ProcessingError("处理进程失败，请检查文件内容或减少文件数量后重试。")
            return _collect_results(output_root, sources, warnings, len(uploads), root)
    except ProcessingError:
        raise
    except (OSError, ValueError, csv.Error, zipfile.BadZipFile, zlib.error) as exc:
        raise ProcessingError("文件处理失败，请检查上传文件是否完整后重试。") from exc


def make_demo_uploads() -> list[Upload]:
    """Build small, explicitly synthetic examples with a 0.005 percent step."""
    points: list[tuple[float, float]] = []
    for index in range(2401):
        strain = index / 200.0
        if strain <= 0.25:
            stress = 2000.0 * strain  # 200 GPa in MPa per percent
        elif strain <= 8.0:
            stress = 500.0 + 400.0 * (strain - 0.25) / 7.75
        else:
            stress = 900.0 - 30.0 * (strain - 8.0)
        points.append((strain, stress))
    points.extend([(12.005, 1.0), (12.01, 0.0), (12.015, 0.0)])
    workbook = Workbook()
    ws = workbook.active
    ws.title = "合成演示曲线"
    ws.append(["Test Run 1", ""])
    ws.append(["%", "MPa"])
    for strain, stress in points:
        ws.append([strain, stress])
    summary = workbook.create_sheet("合成演示汇总")
    summary.append(["名称", "试样标识", "Rp0.2", "Rm", "At", "E(GPa)"])
    summary.append(["Test Run 1", "C900-合成演示-1", 500, 900, 12, 200])
    xlsx_bytes = BytesIO()
    workbook.save(xlsx_bytes)
    workbook.close()
    csv_text = StringIO(newline="")
    writer = csv.writer(csv_text)
    writer.writerow(["序号", "试样标识", "规定塑性延伸强度", "抗拉强度", "断裂总伸长率", "弹性模量"])
    writer.writerow([1, "C900-合成演示-2", 500, 900, 12, 200])
    writer.writerow([])
    writer.writerow(["时间", "拉伸应变", "应力"])
    writer.writerow(["s", "(%)", "(MPa)"])
    writer.writerows((index / 100.0, strain, stress) for index, (strain, stress) in enumerate(points))
    return [Upload("合成演示曲线.xlsx", xlsx_bytes.getvalue()),
            Upload("合成演示曲线.csv", csv_text.getvalue().encode("utf-8-sig"))]
