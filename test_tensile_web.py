"""Web adapter boundary tests; all curves are synthetic, never experimental data."""
from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from io import BytesIO, StringIO
from pathlib import Path
import csv
import binascii
import importlib
import stat
import struct
import subprocess
import sys
import tracemalloc
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image


def api():
    assert Path(__file__).with_name("tensile_web.py").is_file(), "Missing isolated web adapter"
    return importlib.import_module("tensile_web")


def csv_curve(sample="C991-合成演示-1"):
    stream = StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(["序号", "试样标识"])
    writer.writerow([1, sample])
    writer.writerow(["时间", "拉伸应变", "应力"])
    writer.writerow(["s", "(%)", "(MPa)"])
    for i in range(2401):
        strain = i / 200
        stress = 2000 * strain if strain <= .25 else (
            500 + 400 * (strain - .25) / 7.75 if strain <= 8 else 900 - 30 * (strain - 8))
        writer.writerow([i / 100, strain, stress])
    writer.writerows([(24.01, 12.005, 1), (24.02, 12.01, 0), (24.03, 12.015, 0)])
    return stream.getvalue().encode("utf-8-sig")


def zip_bytes(entries):
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries:
            archive.writestr(name, data)
    return stream.getvalue()


def metadata_xlsx():
    wb = Workbook()
    wb.active.append(["名称", "试样标识", "Rp0.2", "Rm", "At", "E(GPa)"])
    wb.active.append(["Test Run 1", "C992-合成演示-1", 500, 900, 12, 200])
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def test_empty_uploads_fail_in_chinese():
    web = api()
    with pytest.raises(web.ProcessingError, match="上传"):
        web.process_uploads([])


@pytest.mark.parametrize("name", ["../evil.csv", "C:\\evil.csv", "/evil.csv", "a/b.csv", "NUL.csv"])
def test_unsafe_upload_names_are_rejected(name):
    web = api()
    with pytest.raises(web.ProcessingError, match="文件名|路径"):
        web.process_uploads([web.Upload(name, b"x")])


def test_case_insensitive_duplicate_names_are_rejected():
    web = api()
    with pytest.raises(web.ProcessingError, match="冲突|同名"):
        web.process_uploads([web.Upload("a.CSV", b"x"), web.Upload("A.csv", b"y")])


@pytest.mark.parametrize("name", ["../evil.csv", "/evil.csv", "C:/evil.csv", "a/../../evil.csv", "a\\..\\evil.csv"])
def test_zip_traversal_is_rejected(name):
    web = api()
    with pytest.raises(web.ProcessingError, match="路径"):
        web.process_uploads([web.Upload("batch.zip", zip_bytes([(name, b"x")]))])


def test_zip_symlink_and_nested_archives_are_rejected():
    web = api()
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        info = zipfile.ZipInfo("link.csv")
        info.create_system = 3
        info.external_attr = (stat.S_IFLNK | 0o777) << 16
        archive.writestr(info, "outside")
    for data in [stream.getvalue(), zip_bytes([("inside.ZIP", b"x")])]:
        with pytest.raises(web.ProcessingError, match="符号链接|嵌套"):
            web.process_uploads([web.Upload("batch.zip", data)])


def test_csv_adaptation_collision_is_rejected():
    web = api()
    data = zip_bytes([("a.csv", b"x"), ("_Exports/a.CSV", b"y")])
    with pytest.raises(web.ProcessingError, match="冲突|同名"):
        web.process_uploads([web.Upload("batch.zip", data)])


def test_output_filename_collision_does_not_silently_replace_a_plot():
    web = api()
    with pytest.raises(web.ProcessingError, match="结果.*冲突|冲突.*结果"):
        web.process_uploads([web.Upload("a b.csv", csv_curve("C997-合成演示-1")),
                            web.Upload("a_b.csv", csv_curve("C998-合成演示-1"))])


def test_loose_workbooks_share_instrument_summary_folder():
    web = api()
    wb = Workbook()
    ws = wb.active
    ws.append(["Test Run 1", ""])
    ws.append(["%", "MPa"])
    for row in list(csv.reader(StringIO(csv_curve().decode("utf-8-sig"))))[4:]:
        ws.append([float(row[1]), float(row[2])])
    content = BytesIO()
    wb.save(content)
    result = web.process_uploads([web.Upload("curve.XLSX", content.getvalue()),
                                  web.Upload("summary.xlsx", metadata_xlsx())])
    assert result.rows[0]["sample_id"] == "C992-合成演示-1"
    assert result.rows[0]["instrument_e_gpa"] == "200.0"
    assert any("summary.xlsx" in w and "未产生曲线" in w for w in result.warnings)


def test_upload_and_expansion_limits_are_enforced(monkeypatch):
    web = api()
    monkeypatch.setattr(web, "MAX_FILE_BYTES", 16)
    with pytest.raises(web.ProcessingError, match="大小|超过"):
        web.process_uploads([web.Upload("a.csv", b"x" * 17)])
    monkeypatch.setattr(web, "MAX_FILE_BYTES", 10000)
    monkeypatch.setattr(web, "MAX_TOTAL_BYTES", 5)
    with pytest.raises(web.ProcessingError, match="总|超过"):
        web.process_uploads([web.Upload("a.csv", b"123"), web.Upload("b.csv", b"123")])
    monkeypatch.setattr(web, "MAX_TOTAL_BYTES", 10000)
    monkeypatch.setattr(web, "MAX_EXPANDED_BYTES", 100)
    with pytest.raises(web.ProcessingError, match="展开|解压"):
        web.process_uploads([web.Upload("a.zip", zip_bytes([("a.csv", b"x" * 101)]))])


def test_forged_zip_size_does_not_cause_unbounded_inflate_allocation(monkeypatch, tmp_path):
    web = api()
    raw = b"x" * (8 * 1024 * 1024)
    forged = bytearray(zip_bytes([("a.csv", raw)]))
    central = forged.index(b"PK\x01\x02")
    one_byte_crc = binascii.crc32(b"x")
    struct.pack_into("<I", forged, 14, one_byte_crc)
    struct.pack_into("<I", forged, 22, 1)
    struct.pack_into("<I", forged, central + 16, one_byte_crc)
    struct.pack_into("<I", forged, central + 24, 1)
    monkeypatch.setattr(web, "MAX_EXPANDED_BYTES", 100)

    tracemalloc.start()
    try:
        with pytest.raises(web.ProcessingError, match="展开|解压|大小|校验"):
            web._stage_uploads([web.Upload("small.zip", bytes(forged))], tmp_path)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()

    assert peak < 2 * 1024 * 1024


def test_xlsx_internal_expansion_is_checked(monkeypatch):
    web = api()
    monkeypatch.setattr(web, "MAX_EXPANDED_BYTES", 100)
    with pytest.raises(web.ProcessingError, match="展开|解压"):
        web.process_uploads([web.Upload("a.xlsx", zip_bytes([("xl/test.xml", b"x" * 101)]))])


def test_data_file_limit_is_checked(monkeypatch):
    web = api()
    monkeypatch.setattr(web, "MAX_DATA_FILES", 2)
    with pytest.raises(web.ProcessingError, match="文件|数量"):
        web.process_uploads([web.Upload("batch.zip", zip_bytes([(f"{i}.csv", b"x") for i in range(3)]))])


def test_invalid_archive_is_an_actionable_error():
    web = api()
    with pytest.raises(web.ProcessingError, match="ZIP|压缩"):
        web.process_uploads([web.Upload("bad.zip", b"not a zip")])


def test_corrupt_deflate_is_an_actionable_error():
    web = api()
    damaged = bytearray(zip_bytes([("bad.csv", b"hello")]))
    name_length, extra_length = struct.unpack_from("<HH", damaged, 26)
    damaged[30 + name_length + extra_length] = 0x07
    with pytest.raises(web.ProcessingError, match="ZIP|压缩|损坏"):
        web.process_uploads([web.Upload("broken.zip", bytes(damaged))])


def test_zero_curve_data_reports_chinese_failure():
    web = api()
    with pytest.raises(web.ProcessingError, match="曲线"):
        web.process_uploads([web.Upload("metadata.xlsx", metadata_xlsx())])


def test_partial_success_reports_bad_and_no_curve_files():
    web = api()
    result = web.process_uploads([
        web.Upload("curve.CSV", csv_curve()),
        web.Upload("broken.xlsx", b"not an xlsx"),
        web.Upload("metadata.xlsx", metadata_xlsx()),
    ])
    assert len(result.rows) == 1
    assert result.uploaded_count == 3
    assert any(e["source_file"].endswith("broken.xlsx") for e in result.errors)
    assert any("metadata.xlsx" in w and "未产生曲线" in w for w in result.warnings)
    assert not any("broken.xlsx" in w and "未产生曲线" in w for w in result.warnings)


def test_demo_returns_readable_output_only_and_relative_picture_paths():
    web = api()
    uploads = web.make_demo_uploads()
    assert {Path(u.name).suffix.lower() for u in uploads} == {".xlsx", ".csv"}
    assert all("演示" in u.name for u in uploads)
    result = web.process_uploads(uploads)
    assert len(result.rows) == 2
    assert {r["source_type"] for r in result.rows} == {"csv", "xlsx"}
    assert all(float(r["slip_count"]) == 0 for r in result.rows)
    assert all(float(r["point_count"]) > 1000 for r in result.rows)
    assert not result.errors
    assert not result.warnings
    for row in result.rows:
        for field in ["file_plot", "group_plot"]:
            assert row[field] in result.files
            assert not Path(row[field]).is_absolute()
    index = list(csv.DictReader(StringIO(result.files["curve_index.csv"].decode("utf-8-sig"))))
    assert index == result.rows
    assert "tensile_web_" not in result.files["curve_index.csv"].decode("utf-8-sig")
    with zipfile.ZipFile(BytesIO(result.archive)) as archive:
        assert set(archive.namelist()) == set(result.files)
        assert not any(Path(u.name).name in archive.namelist() for u in uploads)
    for name, data in result.files.items():
        if name.endswith(".png"):
            with Image.open(BytesIO(data)) as image:
                image.verify()
        elif name.endswith(".xlsx"):
            wb = load_workbook(BytesIO(data), data_only=True)
            assert wb.active.max_row >= 2
            wb.close()


def test_demo_exposes_bounded_curve_data_for_custom_plotting():
    web = api()
    result = web.process_uploads(web.make_demo_uploads())

    assert len(result.plot_curves) == 2
    assert len({curve["curve_id"] for curve in result.plot_curves}) == 2
    for curve in result.plot_curves:
        assert curve["sample_id"].startswith("C900-")
        assert 2 <= len(curve["x"]) == len(curve["y"]) <= 2000
        assert all(isinstance(value, (int, float)) for value in curve["x"])
        assert all(isinstance(value, (int, float)) for value in curve["y"])
    assert "plot_curves.json" not in result.files


def test_overlay_mode_returns_one_png_containing_multiple_curve_colors():
    web = api()
    result = web.process_uploads(web.make_demo_uploads())

    outputs = web.build_plot_outputs(
        result.plot_curves,
        mode="overlay",
        title="Tensile stress-strain curves",
    )

    assert list(outputs) == ["selected_curves.png"]
    with Image.open(BytesIO(outputs["selected_curves.png"])) as image:
        image.verify()
    with Image.open(BytesIO(outputs["selected_curves.png"])) as image:
        assert image.info["dpi"][0] == pytest.approx(300, abs=1)
        saturated_colors = {
            pixel for pixel in image.convert("RGB").getdata()
            if max(pixel) - min(pixel) >= 60 and max(pixel) < 245
        }
        assert len(saturated_colors) >= 2


def test_separate_mode_returns_one_png_per_selected_curve():
    web = api()
    result = web.process_uploads(web.make_demo_uploads())

    outputs = web.build_plot_outputs(result.plot_curves, mode="separate", title="")

    assert len(outputs) == 2
    assert len({name.casefold() for name in outputs}) == 2
    assert all(name.endswith(".png") for name in outputs)
    for data in outputs.values():
        with Image.open(BytesIO(data)) as image:
            image.verify()


def test_custom_plotting_requires_at_least_one_curve():
    web = api()
    with pytest.raises(web.ProcessingError, match="曲线"):
        web.build_plot_outputs([], mode="overlay", title="")


def test_separate_mode_limits_one_request_to_24_high_resolution_images():
    web = api()
    curves = [
        {
            "curve_id": f"curve_{index:04d}",
            "sample_id": f"sample-{index}",
            "run_name": "",
            "x": [0.0, 1.0],
            "y": [0.0, 100.0],
        }
        for index in range(25)
    ]
    with pytest.raises(web.ProcessingError, match="24"):
        web.build_plot_outputs(curves, mode="separate", title="")


def test_two_archives_keep_same_named_curves_separate():
    web = api()
    a = zip_bytes([("folder/curve.CSV", csv_curve("C993-合成演示-1"))])
    b = zip_bytes([("folder/curve.csv", csv_curve("C994-合成演示-1"))])
    result = web.process_uploads([web.Upload("one.zip", a), web.Upload("two.ZIP", b)])
    assert {r["sample_id"] for r in result.rows} == {"C993-合成演示-1", "C994-合成演示-1"}
    assert len({r["source_file"] for r in result.rows}) == 2
    assert all("folder/_Exports/" in r["source_file"] for r in result.rows)


def test_formula_text_is_safe_in_csv_and_xlsx():
    web = api()
    result = web.process_uploads([web.Upload("formula.csv", csv_curve("=2+3"))])
    assert result.rows[0]["sample_id"] == "'=2+3"
    for name, data in result.files.items():
        if name.endswith(".xlsx"):
            wb = load_workbook(BytesIO(data), data_only=False)
            cell = wb.active["A2"]
            assert cell.data_type == "s"
            assert cell.value == "'=2+3"
            wb.close()


def test_parallel_requests_do_not_mix_results():
    web = api()
    def run(sample):
        return web.process_uploads([web.Upload("same.csv", csv_curve(sample))])
    with ThreadPoolExecutor(max_workers=2) as pool:
        first, second = list(pool.map(run, ["C995-合成演示-1", "C996-合成演示-1"]))
    assert [r["sample_id"] for r in first.rows] == ["C995-合成演示-1"]
    assert [r["sample_id"] for r in second.rows] == ["C996-合成演示-1"]


def test_timeout_is_chinese_and_temporary_folder_is_removed(monkeypatch):
    web = api()
    roots = []
    original = web.tempfile.TemporaryDirectory
    def track(*args, **kwargs):
        folder = original(*args, **kwargs)
        roots.append(Path(folder.name))
        return folder
    monkeypatch.setattr(web.tempfile, "TemporaryDirectory", track)
    with pytest.raises(web.ProcessingError, match="超时"):
        web.process_uploads([web.Upload("curve.csv", csv_curve())], timeout_seconds=.001)
    assert roots and all(not path.exists() for path in roots)


def test_subprocess_failure_is_chinese_and_cleaned(monkeypatch, tmp_path):
    web = api()
    roots = []
    original = web.tempfile.TemporaryDirectory
    def track(*args, **kwargs):
        folder = original(*args, **kwargs)
        roots.append(Path(folder.name))
        return folder
    monkeypatch.setattr(web.tempfile, "TemporaryDirectory", track)
    broken = tmp_path / "broken_cli.py"
    broken.write_text("raise RuntimeError('controlled failure')", encoding="utf-8")
    monkeypatch.setattr(web, "_CLI_SCRIPT", broken)
    with pytest.raises(web.ProcessingError, match="处理进程失败"):
        web.process_uploads([web.Upload("curve.csv", csv_curve())])
    assert roots and all(not path.exists() for path in roots)


def test_success_removes_temporary_input_and_output_folders(monkeypatch):
    web = api()
    roots = []
    original = web.tempfile.TemporaryDirectory
    def track(*args, **kwargs):
        folder = original(*args, **kwargs)
        roots.append(Path(folder.name))
        return folder
    monkeypatch.setattr(web.tempfile, "TemporaryDirectory", track)
    result = web.process_uploads([web.Upload("curve.csv", csv_curve())])
    assert result.rows and result.archive
    assert roots and all(not path.exists() for path in roots)


def test_adapter_matches_original_cli_properties(tmp_path):
    web = api()
    source = tmp_path / "input" / "_Exports"
    source.mkdir(parents=True)
    (source / "curve.csv").write_bytes(csv_curve())
    output = tmp_path / "output"
    completed = subprocess.run([sys.executable, str(Path(__file__).with_name("处理拉伸数据.py")),
        "--input", str(source.parent), "--output", str(output)], capture_output=True, timeout=120)
    assert completed.returncode == 0, completed.stderr.decode("utf-8", errors="replace")
    with (output / "curve_index.csv").open(encoding="utf-8-sig", newline="") as stream:
        baseline = list(csv.DictReader(stream))[0]
    actual = web.process_uploads([web.Upload("curve.csv", csv_curve())]).rows[0]
    for field in ["rp02_mpa", "rm_mpa", "at_percent", "e_gpa", "point_count", "slip_count"]:
        assert actual[field] == baseline[field]
